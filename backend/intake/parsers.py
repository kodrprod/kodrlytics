"""Document parsers: PDF, Excel, CSV, SEC EDGAR JSON -> raw text / dict for LLM extraction."""
from __future__ import annotations
import csv
import io
import logging
from datetime import date
from pathlib import Path
from typing import Union

log = logging.getLogger(__name__)


def parse_csv(source: Union[str, bytes, Path]) -> str:
    if isinstance(source, Path):
        text = source.read_text(encoding="utf-8-sig")
    elif isinstance(source, bytes):
        text = source.decode("utf-8-sig")
    else:
        text = source
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    return "\n".join("\t".join(row) for row in rows)


def parse_excel(source: Union[bytes, Path]) -> str:
    import openpyxl
    if isinstance(source, Path):
        wb = openpyxl.load_workbook(source, data_only=True)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(source), data_only=True)
    parts = []
    for name in wb.sheetnames:
        ws = wb[name]
        parts.append(f"=== Sheet: {name} ===")
        for row in ws.iter_rows(values_only=True):
            parts.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def parse_pdf(source: Union[bytes, Path]) -> str:
    import pdfplumber
    if isinstance(source, Path):
        pdf = pdfplumber.open(source)
    else:
        pdf = pdfplumber.open(io.BytesIO(source))
    pages = []
    with pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return "\n\n".join(pages)


def parse_sec_edgar_json(data: dict) -> dict:
    """
    Parse SEC EDGAR company facts JSON (from EDGAR /api/xbrl/companyfacts endpoint)
    into the canonical CompanyFinancials dict format.

    Input format:
      { "cik": "...", "entityName": "Apple Inc.", "facts": { "us-gaap": { "Revenues": {...}, ... } } }
    """
    entity_name = data.get("entityName", "Unknown")
    us_gaap = data.get("facts", {}).get("us-gaap", {})
    log.info("SEC EDGAR parser: entity=%s, concepts available=%d", entity_name, len(us_gaap))

    def _get_flow_values(concepts: list[str]) -> dict[str, float]:
        """Income statement: entries that cover ~12 months (have start + end dates)."""
        for concept in concepts:
            if concept not in us_gaap:
                continue
            usd_data = us_gaap[concept].get("units", {}).get("USD", [])
            annual: dict[str, float] = {}
            for e in usd_data:
                if e.get("form") not in ("10-K", "10-K/A"):
                    continue
                start_s, end_s = e.get("start", ""), e.get("end", "")
                if not start_s or not end_s:
                    continue
                try:
                    s = date.fromisoformat(start_s)
                    end = date.fromisoformat(end_s)
                    months = (end.year - s.year) * 12 + (end.month - s.month)
                    if not (10 <= months <= 14):
                        continue
                except ValueError:
                    continue
                year = end_s[:4]
                val = e.get("val")
                if val is not None:
                    annual[year] = float(val)
            if annual:
                log.debug("  [OK] flow concept %s: %d periods", concept, len(annual))
                return annual
            log.debug("  [--] flow concept %s: no annual data", concept)
        return {}

    def _get_point_values(concepts: list[str]) -> dict[str, float]:
        """Balance sheet: point-in-time entries (no start date)."""
        for concept in concepts:
            if concept not in us_gaap:
                continue
            usd_data = us_gaap[concept].get("units", {}).get("USD", [])
            annual: dict[str, float] = {}
            for e in usd_data:
                if e.get("form") not in ("10-K", "10-K/A"):
                    continue
                if e.get("start"):  # balance sheet = point-in-time, skip period entries
                    continue
                end_s = e.get("end", "")
                if not end_s:
                    continue
                year = end_s[:4]
                val = e.get("val")
                if val is not None:
                    annual[year] = float(val)
            if annual:
                log.debug("  [OK] point concept %s: %d periods", concept, len(annual))
                return annual
            log.debug("  [--] point concept %s: no annual data", concept)
        return {}

    # ── Income statement ───────────────────────────────────────────────────────
    revenue = _get_flow_values([
        "Revenues",
        "RevenueFromContractWithCustomerExcludingAssessedTax",
        "RevenueFromContractWithCustomerIncludingAssessedTax",
        "SalesRevenueNet", "SalesRevenueGoodsNet",
    ])
    if not revenue:
        raise ValueError("No annual revenue data found in SEC EDGAR JSON")

    # Use last 4 fiscal years
    periods = sorted(sorted(revenue.keys(), reverse=True)[:4])
    log.info("SEC EDGAR parser: using periods %s", periods)

    def _f(d: dict) -> dict[str, float]:
        return {y: d[y] for y in periods if y in d}

    rev   = _f(revenue)
    cogs  = _f(_get_flow_values(["CostOfRevenue", "CostOfGoodsSold", "CostOfGoodsSoldAndServicesCosts"]))
    gross = _f(_get_flow_values(["GrossProfit"]))
    opex  = _f(_get_flow_values(["OperatingExpenses", "SellingGeneralAndAdministrativeExpense"]))
    ebit  = _f(_get_flow_values(["OperatingIncomeLoss"]))
    inter = _f(_get_flow_values(["InterestExpense", "InterestAndDebtExpense", "InterestExpenseDebt"]))
    ebt   = _f(_get_flow_values([
        "IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest",
        "IncomeLossFromContinuingOperationsBeforeIncomeTaxesMinorityInterestAndIncomeLossFromEquityMethodInvestments",
    ]))
    tax    = _f(_get_flow_values(["IncomeTaxExpenseBenefit"]))
    net_in = _f(_get_flow_values(["NetIncomeLoss"]))
    da     = _f(_get_flow_values(["DepreciationDepletionAndAmortization", "DepreciationAndAmortization"]))

    # ── Balance sheet ──────────────────────────────────────────────────────────
    cash   = _f(_get_point_values(["CashAndCashEquivalentsAtCarryingValue", "Cash",
                                    "CashAndShortTermInvestments"]))
    ar     = _f(_get_point_values(["AccountsReceivableNetCurrent", "ReceivablesNetCurrent",
                                    "NontradeReceivablesCurrent"]))
    inv    = _f(_get_point_values(["InventoryNet"]))
    cur_a  = _f(_get_point_values(["AssetsCurrent"]))
    ppe    = _f(_get_point_values(["PropertyPlantAndEquipmentNet"]))
    tot_a  = _f(_get_point_values(["Assets"]))
    ap     = _f(_get_point_values(["AccountsPayableCurrent"]))
    st_dbt = _f(_get_point_values(["LongTermDebtCurrent", "ShortTermBorrowings",
                                    "CommercialPaper"]))
    cur_l  = _f(_get_point_values(["LiabilitiesCurrent"]))
    lt_dbt = _f(_get_point_values(["LongTermDebtNoncurrent", "LongTermDebt",
                                    "LongTermNotesPayable"]))
    tot_l  = _f(_get_point_values(["Liabilities"]))
    ret_e  = _f(_get_point_values(["RetainedEarningsAccumulatedDeficit"]))
    tot_eq = _f(_get_point_values([
        "StockholdersEquity", "StockholdersEquityAttributableToParent", "Equity",
    ]))

    # ── Derive missing items ───────────────────────────────────────────────────
    for y in periods:
        if y not in gross and y in rev and y in cogs:
            gross[y] = rev[y] - cogs[y]
        if y not in ebt and y in net_in and y in tax:
            ebt[y] = net_in[y] + tax[y]
        if y not in tot_eq and y in tot_a and y in tot_l:
            tot_eq[y] = tot_a[y] - tot_l[y]

    other_ca = {y: max(0.0, cur_a.get(y, 0) - cash.get(y, 0) - ar.get(y, 0) - inv.get(y, 0))
                for y in periods if y in cur_a}
    other_cl = {y: max(0.0, cur_l.get(y, 0) - ap.get(y, 0) - st_dbt.get(y, 0))
                for y in periods if y in cur_l}
    other_nca = {y: max(0.0, tot_a.get(y, 0) - cur_a.get(y, 0) - ppe.get(y, 0))
                 for y in periods if y in tot_a and y in cur_a}
    other_ncl = {y: max(0.0, tot_l.get(y, 0) - cur_l.get(y, 0) - lt_dbt.get(y, 0))
                 for y in periods if y in tot_l and y in cur_l}

    # Log what we found vs missed
    found = {k: bool(v) for k, v in [
        ("revenue", rev), ("cogs", cogs), ("gross_profit", gross), ("ebit", ebit),
        ("net_income", net_in), ("total_assets", tot_a), ("total_equity", tot_eq),
        ("total_liabilities", tot_l),
    ]}
    log.info("SEC EDGAR extraction: %s",
             " | ".join(f"{'[OK]' if ok else '[--]'} {k}" for k, ok in found.items()))

    return {
        "company_name": entity_name,
        "currency": "USD",
        "nace_code": "J",  # Information/Technology default for SEC tech filers
        "reporting_standard": "US-GAAP",
        "income_statement": {
            "periods": periods,
            "revenue": rev, "cost_of_goods_sold": cogs, "gross_profit": gross,
            "operating_expenses": opex, "ebit": ebit, "interest_expense": inter,
            "ebt": ebt, "income_tax": tax, "net_income": net_in,
            "depreciation_amortization": da,
        },
        "balance_sheet": {
            "periods": periods,
            "cash": cash, "accounts_receivable": ar, "inventory": inv,
            "other_current_assets": other_ca, "current_assets": cur_a,
            "fixed_assets": ppe, "other_noncurrent_assets": other_nca,
            "total_assets": tot_a,
            "accounts_payable": ap, "short_term_debt": st_dbt,
            "other_current_liabilities": other_cl, "current_liabilities": cur_l,
            "long_term_debt": lt_dbt, "other_noncurrent_liabilities": other_ncl,
            "total_liabilities": tot_l,
            "share_capital": {}, "retained_earnings": ret_e, "total_equity": tot_eq,
        },
    }


def parse_document(filename: str, content: bytes) -> str:
    """Dispatch to the right parser based on file extension."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(content)
    elif ext in (".xlsx", ".xls"):
        return parse_excel(content)
    elif ext == ".csv":
        return parse_csv(content)
    else:
        return content.decode("utf-8", errors="replace")
